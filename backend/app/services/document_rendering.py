"""Renders the structured content dict built by DocumentGeneratorService
into each export format. Markdown is the source of truth; HTML is
derived from it. DOCX and PDF are built directly from the content dict
so their structure (headings, tables) doesn't depend on markdown-to-X
conversion fidelity.
"""
import io

import markdown2
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def render_markdown(content: dict) -> str:
    if "cahier_des_charges" in content:
        return render_business_markdown(content)

    p = content["project"]
    lines = [
        f"# Cahier des charges — {p['name']}",
        "",
        f"*{p['organization']} — généré le {content['generated_at']}*",
        "",
        "## 1. Résumé exécutif",
        "",
        content["executive_summary"],
        "",
        f"**Statut du projet :** {p['status']} · "
        f"**Complétude du cadrage :** {content['completeness']['overall_completion']}% · "
        f"**Confiance moyenne :** {content['completeness']['overall_confidence']}%",
        "",
    ]

    if content["stakeholders"]:
        lines += ["## 2. Parties prenantes", ""]
        for s in content["stakeholders"]:
            lines.append(f"- **{s['concept']}** : {s['answer']}")
        lines.append("")

    lines += ["## 3. Exigences", ""]
    for req_type, reqs in content["requirements_by_type"].items():
        label = content["requirement_type_labels"].get(req_type, req_type)
        lines.append(f"### {label}")
        lines.append("")
        for r in reqs:
            lines.append(f"**{r['requirement_key']} — {r['title']}** (priorité : {r['priority']}, risque : {r['risk']})")
            lines.append("")
            lines.append(r["description"])
            lines.append("")
            if r["acceptance_criteria"]:
                lines.append("Critères d'acceptation :")
                for c in r["acceptance_criteria"]:
                    lines.append(f"- {c}")
                lines.append("")
            if r["dependencies"]:
                lines.append(f"Dépendances : {', '.join(r['dependencies'])}")
                lines.append("")
    if not content["requirements_by_type"]:
        lines.append("_Aucune exigence n'a encore été générée pour ce projet._")
        lines.append("")

    lines += ["## 4. Dictionnaire de données", ""]
    if content["data_dictionary"]:
        lines.append("| Nom | Domaine | Description |")
        lines.append("|---|---|---|")
        for d in content["data_dictionary"]:
            lines.append(f"| {d['name']} | {d['domain']} | {d['description']} |")
        lines.append("")
    else:
        lines.append("_Aucune entité documentée pour le moment._")
        lines.append("")

    lines += ["## 5. User Stories", ""]
    for us in content["user_stories"]:
        lines.append(f"- **{us['requirement_key']}** — {us['story']}")
    if not content["user_stories"]:
        lines.append("_Aucune user story disponible._")
    lines.append("")

    lines += ["## 6. Plan de test", ""]
    if content["test_plan"]:
        lines.append("| Exigence | Cas de test | Priorité | Statut |")
        lines.append("|---|---|---|---|")
        for t in content["test_plan"]:
            lines.append(f"| {t['requirement_key']} | {t['test_case']} | {t['priority']} | {t['status']} |")
        lines.append("")
    else:
        lines.append("_Aucun cas de test disponible._")
        lines.append("")

    lines += ["## 7. Conception — Diagramme MVP", ""]
    lines.append(
        "Diagramme unique de conception représentant le MVP : acteurs, frontend, "
        "backend, modules justifiés par les exigences, données et contrôles de "
        "sécurité. Conformément au principe de conception de ce document, seuls "
        "les éléments requis par les exigences confirmées apparaissent — pas de "
        "briques ajoutées par défaut (microservices, Kubernetes, etc.)."
    )
    lines.append("")
    mvp_diagram = content["diagrams"].get("mvp_diagram_mermaid", "")
    if mvp_diagram:
        lines += ["```mermaid", mvp_diagram, "```", ""]

    sec = content["security"]
    lines += ["## 8. Sécurité", ""]
    if not sec["available"]:
        lines.append(sec["note"])
        lines.append("")
    else:
        lines.append(f"**Score de sécurité global : {sec['security_score']}/100**")
        lines.append("")
        lines.append("| Contrôle | Statut | Recommandation |")
        lines.append("|---|---|---|")
        for c in sec["checklist"]:
            lines.append(f"| {c['name']} | {c['status']} | {c['recommendation']} |")
        lines.append("")
        lines.append("### Registre des risques")
        lines.append("")
        for r in sec["risk_register"]:
            lines.append(f"- **{r['risk_id']} — {r['title']}** (impact : {r['impact']}, statut : {r['status']}) — {r['mitigation']}")
        lines.append("")

    lines += ["## 9. Traçabilité", ""]
    lines.append(
        "Chaque exigence référence le concept du Knowledge Graph dont elle est issue "
        "(`source_concept_key`) — voir `/requirements/traceability-matrix` pour la matrice complète."
    )
    lines.append("")

    lines += ["## 10. Déploiement, maintenance et formation", ""]
    lines.append(f"**Déploiement :** {content['deployment_notes'] or 'Non documenté.'}")
    lines.append("")
    lines.append(f"**Disponibilité :** {content['availability_notes'] or 'Non documentée.'}")
    lines.append("")
    lines.append(f"**Maintenance / recette :** {content['maintenance_notes'] or 'Non documentée.'}")
    lines.append("")
    lines.append(f"**Formation :** {content['training_notes'] or 'Non documentée.'}")
    lines.append("")
    lines.append(f"**Supervision :** {content['monitoring_notes'] or 'Non documentée.'}")
    lines.append("")

    if content["unimplemented_sections"]:
        lines += ["## 11. Sections non générées", ""]
        for s in content["unimplemented_sections"]:
            lines.append(f"- **{s['name']}** — {s['reason']}")
        lines.append("")

    return "\n".join(lines)


def render_business_markdown(content: dict) -> str:
    p = content["project"]
    cdc = content["cahier_des_charges"]
    date_str = content.get("generated_at", "")[:10]
    org_str = p.get("organization") or "Organisation Client"

    lines = [
        f"# Cahier des Charges — {p['name']}",
        "",
        f"**Organisation :** {org_str}  ",
        f"**Date :** {date_str} | **Statut :** {p.get('status', 'Brouillon')} | **Complétude :** {cdc['completeness_score']}%",
        "",
        "---",
        "",
        "## RÉSUMÉ ÉXÉCUTIF",
        "",
        cdc["summary"],
        "",
    ]

    for section in cdc["sections"]:
        lines += [f"## {section['title']}", ""]
        for item in section["items"]:
            if item.startswith("•") or item.startswith("-"):
                lines.append(item)
            else:
                lines.append(f"• {item}")
        lines.append("")

    if cdc.get("points_to_confirm"):
        lines += ["## POINTS À VALIDER", ""]
        for point in cdc["points_to_confirm"]:
            lines.append(f"- [ ] {point}")
        lines.append("")

    return "\n".join(lines)


def render_latex(content: dict) -> str:
    md = render_markdown(content)
    escaped = (
        md.replace("\\", "\\textbackslash{}")
        .replace("&", "\\&")
        .replace("%", "\\%")
        .replace("$", "\\$")
        .replace("#", "\\#")
        .replace("_", "\\_")
        .replace("{", "\\{")
        .replace("}", "\\}")
    )
    body = []
    for line in escaped.splitlines():
        if line.startswith("\\#\\# "):
            body.append(f"\\section*{{{line[6:]}}}")
        elif line.startswith("\\#\\#\\# "):
            body.append(f"\\subsection*{{{line[8:]}}}")
        elif line.startswith("\\# "):
            body.append(f"\\title{{{line[3:]}}}\\maketitle")
        elif line.startswith("- "):
            body.append(f"\\noindent\\textbullet\\ {line[2:]}\\\\")
        elif line.startswith("```"):
            continue
        else:
            body.append(line + "\n")
    return "\\documentclass[11pt]{article}\n\\usepackage[utf8]{inputenc}\n\\begin{document}\n" + "\n".join(body) + "\n\\end{document}\n"


def render_html(markdown_content: str) -> str:
    body = markdown2.markdown(
        markdown_content, extras=["tables", "fenced-code-blocks", "header-ids"]
    )
    return f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8">
<title>Cahier des charges</title>
<style>
body {{ font-family: -apple-system, Segoe UI, sans-serif; max-width: 900px; margin: 2rem auto; padding: 0 1rem; line-height: 1.6; color: #1a1a1a; }}
table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; }}
th, td {{ border: 1px solid #ddd; padding: 0.5rem; text-align: left; font-size: 0.9rem; }}
th {{ background: #f4f4f4; }}
code {{ background: #f4f4f4; padding: 0.1rem 0.3rem; border-radius: 3px; }}
pre {{ background: #f4f4f4; padding: 1rem; overflow-x: auto; border-radius: 6px; }}
h1, h2, h3 {{ color: #0284c7; }}
</style>
</head>
<body>
{body}
</body>
</html>"""


def render_docx(content: dict) -> bytes:

    return render_business_docx(content)


def render_business_docx(content: dict) -> bytes:

    doc = Document()
    p = content["project"]
    cdc = content["cahier_des_charges"]
    date_str = content.get("generated_at", "")[:10]

    doc.add_heading(f"Cahier des Charges — {p['name']}", level=0)
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Organisation : {p.get('organization') or 'Client'} | Date : {date_str} | Complétude : {cdc['completeness_score']}%").italic = True

    doc.add_heading("Résumé Exécutif", level=1)
    doc.add_paragraph(cdc["summary"])

    for section in cdc["sections"]:
        doc.add_heading(section["title"], level=1)
        for item in section["items"]:
            clean_item = item.lstrip("•- ").strip()
            doc.add_paragraph(clean_item, style="List Bullet")

    if cdc.get("points_to_confirm"):
        doc.add_heading("Points à Valider", level=1)
        for point in cdc["points_to_confirm"]:
            doc.add_paragraph(point, style="List Bullet")

    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(content: dict) -> bytes:
    if "cahier_des_charges" in content:
        return render_business_pdf(content)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm, leftMargin=2 * cm, rightMargin=2 * cm
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], textColor=colors.HexColor("#14532d"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], textColor=colors.HexColor("#14532d"))
    body = styles["BodyText"]

    p = content["project"]
    story = [
        Paragraph(f"Cahier des charges — {p['name']}", styles["Title"]),
        Paragraph(f"{p['organization']} — généré le {content['generated_at']}", body),
        Spacer(1, 12),
        Paragraph("1. Résumé exécutif", h1),
        Paragraph(content["executive_summary"] or "Non documenté.", body),
        Spacer(1, 12),
    ]

    story.append(Paragraph("2. Exigences", h1))
    if not content["requirements_by_type"]:
        story.append(Paragraph("Aucune exigence n'a encore été générée pour ce projet.", body))
    for req_type, reqs in content["requirements_by_type"].items():
        label = content["requirement_type_labels"].get(req_type, req_type)
        story.append(Paragraph(label, h2))
        for r in reqs:
            story.append(
                Paragraph(
                    f"<b>{r['requirement_key']} — {r['title']}</b> (priorité : {r['priority']}, risque : {r['risk']})",
                    body,
                )
            )
            story.append(Paragraph(r["description"] or "", body))
            story.append(Spacer(1, 6))

    story.append(Paragraph("3. Dictionnaire de données", h1))
    if content["data_dictionary"]:
        data = [["Nom", "Domaine", "Description"]] + [
            [d["name"], d["domain"], d["description"][:120]] for d in content["data_dictionary"]
        ]
        table = Table(data, colWidths=[4 * cm, 3 * cm, 9 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14532d")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("Aucune entité documentée pour le moment.", body))

    sec = content["security"]
    story.append(Spacer(1, 12))
    story.append(Paragraph("4. Sécurité", h1))
    if not sec["available"]:
        story.append(Paragraph(sec["note"], body))
    else:
        story.append(Paragraph(f"Score de sécurité global : {sec['security_score']}/100", body))
        data = [["Contrôle", "Statut", "Recommandation"]] + [
            [c["name"], c["status"], c["recommendation"][:100]] for c in sec["checklist"]
        ]
        table = Table(data, colWidths=[4 * cm, 2.5 * cm, 9.5 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14532d")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table)

    doc.build(story)
    return buf.getvalue()


def render_business_pdf(content: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
    )
    styles = getSampleStyleSheet()

    doc_title = ParagraphStyle(
        "CDCTitle", parent=styles["Title"], fontSize=16, leading=20, textColor=colors.HexColor("#0f172a"), alignment=0, spaceAfter=8
    )
    meta_style = ParagraphStyle(
        "CDCMeta", parent=styles["BodyText"], fontSize=9, leading=12, textColor=colors.HexColor("#475569"), spaceAfter=14
    )
    h1 = ParagraphStyle(
        "CDCH1", parent=styles["Heading1"], fontSize=11, leading=14, textColor=colors.HexColor("#0284c7"), spaceBefore=10, spaceAfter=4, keepWithNext=True
    )
    body = ParagraphStyle(
        "CDCBody", parent=styles["BodyText"], fontSize=9, leading=12.5, textColor=colors.HexColor("#1e293b"), spaceAfter=3
    )
    bullet = ParagraphStyle(
        "CDCBullet", parent=styles["BodyText"], fontSize=9, leading=12.5, textColor=colors.HexColor("#1e293b"), leftIndent=12, spaceAfter=3
    )
    confirm_style = ParagraphStyle(
        "CDCConfirm", parent=styles["BodyText"], fontSize=9, leading=12.5, textColor=colors.HexColor("#b45309"), leftIndent=12, spaceAfter=3
    )

    p = content["project"]
    cdc = content["cahier_des_charges"]
    date_str = content.get("generated_at", "")[:10]
    org_str = p.get("organization") or "Organisation Client"

    story = [
        Paragraph(f"Cahier des Charges — {p['name']}", doc_title),
        Paragraph(f"<b>Organisation :</b> {org_str} &nbsp;|&nbsp; <b>Date :</b> {date_str} &nbsp;|&nbsp; <b>Complétude :</b> {cdc['completeness_score']}%", meta_style),
        Spacer(1, 4),
        Paragraph("RÉSUMÉ ÉXÉCUTIF", h1),
        Paragraph(cdc["summary"] or "Cadrage fonctionnel du projet.", body),
        Spacer(1, 6),
    ]

    for section in cdc["sections"]:
        story.append(Paragraph(section["title"], h1))
        for item in section["items"]:
            clean_item = item.lstrip("•- ").strip()
            story.append(Paragraph(f"• {clean_item}", bullet))
        story.append(Spacer(1, 4))

    if cdc.get("points_to_confirm"):
        story.append(Paragraph("POINTS À VALIDER", h1))
        for point in cdc["points_to_confirm"]:
            story.append(Paragraph(f"⚠️ {point}", confirm_style))

    doc.build(story)
    return buf.getvalue()

